"""Artifact checks for this revision; numerical physics checks run in MATLAB."""
from pathlib import Path
import sys,json,re
root=Path(__file__).resolve().parents[2];sys.path.insert(0,str(root/'.task_review_tools'))
import pymupdf
from pptx import Presentation
from docx import Document
out=root/'Lab1_Deliverables';src=out/'Presentation_Source'
poster=pymupdf.open(out/'Gulo_3000_Poster.pdf');hand=pymupdf.open(out/'Gulo_3000_Handout.pdf')
assert len(poster)==1 and len(hand)==3
pt=' '.join(p.get_text() for p in poster);ht=' '.join(p.get_text() for p in hand)
assert 'rigid' not in pt.lower() and 'rigid' not in ht.lower()
assert '105 mm motor wheel + 8.00 mm shaft' in pt
assert '137 → 105 mm' in ht and '3.18 → 8.00 mm' in ht
assert '137 → 100 mm' not in ht and '3.18 → 6.35 mm' not in ht
assert '60 Hz' in pt and '60' in ht
for v in ['165','4.50','1.6','10,000','49.9']:assert v in ht
assert 'di/dt' not in ht and '∫' not in ht
for p in hand:
 for block in p.get_text('dict')['blocks']:
  if 'lines' not in block:continue
  for line in block['lines']:
   for span in line['spans']:
    x0,y0,x1,y1=span['bbox'];assert x0>=0 and y0>=0 and x1<=p.rect.width+.1 and y1<=p.rect.height+.1
layout=json.loads((src/'handout_layout_audit.json').read_text())
assert all(p['font']==12 for p in layout)
assert not json.loads((src/'poster_text_overflow.json').read_text(encoding='utf-8-sig'))
pv=json.loads((out/'Analysis/physical_verification.json').read_text())
assert pv['allFourNominalTargetsPass'] and pv['fullInductancePeakAlsoPasses']
assert all(v>=30 for v in pv['fineGridRelativeSpeedReductions_pct'])
assert all(s['TransientReduction_pct']>=30 for s in pv['jointComplianceScenario'])
prs=Presentation(out/'Gulo_3000_Poster.pptx')
assert len(prs.slides)==1
doc=Document(out/'Gulo_3000_Handout.docx')
assert sum('w:type="page"' in p._p.xml for p in doc.paragraphs)==2
assert any('137 → 105 mm' in p.text and '3.18 → 8.00 mm' in p.text for p in doc.paragraphs)
result={'posterPages':1,'handoutPages':3,'handoutBodyFont_pt':12,
 'handoutWordCountIncludingReferences':len(re.findall(r'\S+',ht)),
 'matlabSourceFiles':len(list((out/'MATLAB').rglob('*.m'))),'noPosterTextOverflow':True,
 'rigidCurveRemoved':True,'physicalDimensionsConsistent':True,'nominalTargetsPass':True,
 'jointComplianceScenarioMotionTargetsPass':True,'editableSourcesPresent':True}
(out/'Analysis/final_deliverable_checks.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result,indent=2))
