"""Build editable presentation documents from MATLAB-generated results.
No engineering calculations are performed here; only formatting and layout.
Requires python-pptx, python-docx, Pillow and PyMuPDF.
"""
from pathlib import Path
import sys, json, csv

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.task_review_tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = ROOT / 'Lab1_Deliverables'
FIG = OUT / 'Figures'
V = json.loads((OUT / 'Analysis/report_values.json').read_text())
P = V['parameters']
DES = {d['Parameter']: d for d in V['designs']}
NAVY = '062D48'; TEAL = '007E80'; GOLD = 'F3C84B'
GRAY = '526577'; PALE = 'EDF3F6'; RED = 'C92934'; WHITE = 'FFFFFF'

def percent(x): return f'{x:.1f}%'
def color(x): return RGBColor.from_string(x)

prs = Presentation()
prs.slide_width = Inches(24)
prs.slide_height = Inches(13.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(WHITE)

def box(x,y,w,h,fill=WHITE,line=None,width=1):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill)
    if line: sh.line.color.rgb = color(line); sh.line.width = Pt(width)
    else: sh.line.fill.background()
    return sh

def text(x,y,w,h,s,size=19,bold=False,fg=NAVY,align=PP_ALIGN.LEFT,margin=.015):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin)
    tf.margin_top=tf.margin_bottom=Inches(margin)
    for i,line in enumerate(s.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(1)
        p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fg)
        p.line_spacing=1.07
    return sh

def image(name,x,y,w,h=None):
    return slide.shapes.add_picture(str(FIG/(name+'.png')),Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None)

def header(x,y,w,label):
    box(x,y,w,.44,NAVY)
    text(x+.11,y+.04,w-.22,.35,label,21,True,WHITE)

def placeholder(x,y,w,h,title,description):
    box(x,y,w,h,'FFF4F4',RED,2)
    text(x+.14,y+.09,w-.28,.3,title,18,True,RED)
    text(x+.14,y+.44,w-.28,h-.5,description,17,False,RED)

def table(x,y,widths,rows,row_h=.35,size=16):
    for r,row in enumerate(rows):
        xx=x
        for c,(s,w) in enumerate(zip(row,widths)):
            fill=NAVY if r==0 else (PALE if r%2 else WHITE)
            box(xx,y+r*row_h,w,row_h,fill)
            text(xx+.08,y+r*row_h+.035,w-.16,row_h-.04,str(s),size,r==0,WHITE if r==0 else NAVY)
            xx+=w

doc=Document();sec=doc.sections[0]
sec.page_width=DInches(8.5);sec.page_height=DInches(11)
sec.top_margin=DInches(.48);sec.bottom_margin=DInches(.43)
sec.left_margin=sec.right_margin=DInches(.52)
sec.header_distance=DInches(.16);sec.footer_distance=DInches(.17)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=DPt(10)
normal.paragraph_format.space_after=DPt(4);normal.paragraph_format.line_spacing=1.02
for name,sz in [('Title',18),('Heading 1',12),('Heading 2',10)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=DPt(sz);st.font.bold=True;st.font.color.rgb=DColor.from_string(NAVY)
    st.paragraph_format.space_before=DPt(5);st.paragraph_format.space_after=DPt(3)
headerp=sec.header.paragraphs[0];headerp.text='POWERTRAIN DYNAMICS GROUP   |   ME 495   |   SEPTEMBER 2026'
headerp.runs[0].font.size=DPt(7);headerp.runs[0].font.color.rgb=DColor.from_string(GRAY)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Gulo 3000 • Technical handout   |   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
for r in fp.runs:r.font.size=DPt(7)

def para(s,style=None):
    return doc.add_paragraph(s,style)
def head(s):return doc.add_paragraph(s,'Heading 1')
def pic(name,width=7.40):
    p=doc.add_paragraph();p.paragraph_format.space_after=DPt(1)
    p.add_run().add_picture(str(FIG/(name+'.png')),width=DInches(width));return p
def caption(s):
    p=para(s)
    for r in p.runs:r.font.size=DPt(8);r.font.color.rgb=DColor.from_string(GRAY)
    return p
def dtable(rows,widths=None,fontsize=8):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    if widths:
        for col,wd in zip(t.columns,widths):col.width=DInches(wd)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        if widths:
            for cell,wd in zip(cells,widths):cell.width=DInches(wd)
        for c,s in zip(cells,row):
            c.text=str(s)
            for p in c.paragraphs:
                p.paragraph_format.space_after=DPt(2);p.paragraph_format.space_before=DPt(1)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs:
                    r.font.name='Arial';r.font.size=DPt(fontsize);r.font.bold=i==0
                    r.font.color.rgb=DColor.from_string(WHITE if i==0 else NAVY)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),NAVY if i==0 else (PALE if i%2 else WHITE));c._tc.get_or_add_tcPr().append(sh)
        trpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    return t

