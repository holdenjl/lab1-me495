"""Build editable presentation documents from MATLAB-generated results.
No engineering calculations are performed here; only formatting and layout.
Requires python-pptx, python-docx, Pillow and PyMuPDF.
"""
from pathlib import Path
import sys, json, csv

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.task_review_tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = ROOT / 'Lab1_Deliverables'
FIG = OUT / 'Figures'
V = json.loads((OUT / 'Analysis/report_values.json').read_text())
P = V['parameters']
DES = {d['Parameter']: d for d in V['designs']}
NAVY = '062D48'; TEAL = '007E80'; GOLD = 'F3C84B'
GRAY = '526577'; PALE = 'EDF3F6'; RED = 'C92934'; WHITE = 'FFFFFF'

def percent(x): return f'{x:.1f}%'
def color(x): return RGBColor.from_string(x)

prs = Presentation()
prs.slide_width = Inches(24)
prs.slide_height = Inches(13.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(WHITE)

def box(x,y,w,h,fill=WHITE,line=None,width=1):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill)
    if line: sh.line.color.rgb = color(line); sh.line.width = Pt(width)
    else: sh.line.fill.background()
    return sh

def text(x,y,w,h,s,size=19,bold=False,fg=NAVY,align=PP_ALIGN.LEFT,margin=.015):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin)
    tf.margin_top=tf.margin_bottom=Inches(margin)
    for i,line in enumerate(s.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(1)
        p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fg)
        p.line_spacing=1.07
    return sh

def image(name,x,y,w,h=None):
    return slide.shapes.add_picture(str(FIG/(name+'.png')),Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None)

def header(x,y,w,label):
    box(x,y,w,.44,NAVY)
    text(x+.11,y+.04,w-.22,.35,label,21,True,WHITE)

def placeholder(x,y,w,h,title,description):
    box(x,y,w,h,'FFF4F4',RED,2)
    text(x+.14,y+.09,w-.28,.3,title,18,True,RED)
    text(x+.14,y+.44,w-.28,h-.5,description,17,False,RED)

def table(x,y,widths,rows,row_h=.35,size=16):
    for r,row in enumerate(rows):
        xx=x
        for c,(s,w) in enumerate(zip(row,widths)):
            fill=NAVY if r==0 else (PALE if r%2 else WHITE)
            box(xx,y+r*row_h,w,row_h,fill)
            text(xx+.08,y+r*row_h+.035,w-.16,row_h-.04,str(s),size,r==0,WHITE if r==0 else NAVY)
            xx+=w

box(0,0,24,1.45,NAVY)
text(.45,.23,23.1,.67,'Four passive changes can cut modeled vibration by ≥30%',38,True,WHITE)
text(.48,.99,23,.31,'Gulo 3000  |  Experimental model + design simulations  |  Powertrain Dynamics Group  |  ME 495',19,False,WHITE)
box(.45,1.57,23.1,.43,GOLD)
text(.58,1.61,22.85,.35,'First prototype to investigate: a shaft-path damper that preserves modeled steady-speed response.',19,True)

x1,x2,x3,w=.45,8.25,16.05,7.5
header(x1,2.12,w,'1  Measure and model the shaft')
text(x1,2.72,w,.69,'Motor switching excites torsional vibration.\nGoal: validate its source and identify four ≥30% solutions.',19)
placeholder(x1,3.60,w,1.34,'PHOTO TO TAKE — physical test rig',
            'Take a clear front/oblique photo. Label the motor, flywheels 1 and 2, shaft, distal bearing and both tachometers.')
placeholder(x1,5.11,w,1.37,'CAD / SCHEMATIC TO ADD — two-inertia system',
            'Capture a simplified CAD assembly or draw a clean schematic. Label Jc, Jf, K, Bs, Bm, Bb, voltage input Vin and speed output ω2.')
text(x1,6.70,w,.35,'Parameters come from measurements and mechanics',20,True)
table(x1,7.13,[3.3,4.2],[['Parameter','Identified value'],
 ['Motor resistance Rm',f'{P["Rm"]:.3f} Ω'],
 ['Motor constant Km',f'{P["Km"]:.4f} V·s/rad'],
 ['Combined damping Btr / Btl',f'{V["Btr"]:.5f} / {V["Btl"]:.5f} N·m·s/rad'],
 ['Effective shaft stiffness K',f'{V["K_refined"]:.3f} N·m/rad (refined)']],.35,16)
image('damping_evidence',x1,9.02,w,2.88)
text(x1,12.02,w,.55,'3 sweeps + 28 time records. Free decay identifies damping; the uncoupled-motor test estimates external loss.',16)

header(x2,2.12,w,'2  The flexible model explains resonance')
image('model_validation',x2,2.73,w,4.65)
box(x2,7.57,w,.74,PALE)
text(x2+.12,7.61,w-.24,.65,
     f'Resonance: {V["measuredResonance"]:.2f} Hz measured / {V["modelResonance"]:.2f} Hz model\n4–10 Hz fit error: {V["refinedMagnitudeRMSE"]:.2f} dB magnitude; {V["refinedPhaseRMSE"]:.2f}° phase',18,True)
text(x2,8.54,w,.39,'Sensitivity separates attenuation from detuning',20,True)
image('sensitivity_screen',x2,9.03,w,2.79)
text(x2,12.02,w,.59,'All 12 parameters were screened. Diameter strongly shifts the peak; Rm, Km, Bm and Bs most strongly change its height.',16)

header(x3,2.12,w,'3  Four quantified design alternatives')
image('four_designs',x3,2.73,w,4.65)
rows=[['Change','Target value','Gain ↓','Pulse ↓']]
for n,label in [('Rm','Resistance'),('Km','Motor constant'),('Bs','Shaft damping'),('Bm','Motor damping')]:
    d=DES[n]
    target={'Rm':f'{d["Proposed"]:.2f} Ω','Km':f'{d["Proposed"]:.3f} V·s/rad','Bs':f'{d["Proposed"]:.3f} N·m·s/rad','Bm':f'{d["Proposed"]:.3f} N·m·s/rad'}[n]
    rows.append([label,target,percent(d['PeakReduction_pct']),percent(d['TransientReduction_pct'])])
table(x3,7.60,[2.08,2.77,1.3,1.35],rows,.36,15.6)
text(x3,9.54,w,.67,'Shaft damper: no local peak; gain assessed over 4–10 Hz.\nAll four reduce the 1-Hz response cost by ≥30%.',16)
image('drive_response',x3,10.39,w,1.56)
text(x3,12.03,w,.56,'Prioritize the shaft damper. Its large damping increase needs hardware testing; the other options sacrifice steady-speed response.',16,True)

box(.45,12.84,23.1,.035,GOLD)
text(.48,12.99,23.04,.29,
     'Predictions at fixed voltage, not modified-hardware test results. Pulse: +1 V for 4 s, then driven to 0 V. Sources: task letter; Lab 1_1 / 1_2; student guide. See three-page handout.',13.7)

slide.notes_slide.notes_text_frame.text = (
    'Suggested 10-minute delivery: 0:00–1:00 problem and recommendation; 1:00–3:00 measurements and damping; '
    '3:00–5:00 model validation and rigid-shaft comparison; 5:00–6:15 sensitivity and detuning; '
    '6:15–8:30 four alternatives and fixed-voltage tradeoffs; 8:30–10:00 shaft-damper recommendation and limits. '
    'Use the red boxes to add a rig photo and a labeled CAD/schematic before presenting. '
    'A rigid shaft means zero relative motion / infinite stiffness, not zero damping. '
    'The shaft-damping design has no local response peak; the table reports its maximum gain in the measured 4–10 Hz band. '
    'Do not claim that four designs were physically tested, or that vehicle comfort was measured. '
    'At equal final speed, the shaft damper retains its pulse benefit; the others lose much of their advantage.'
)
prs.save(OUT/'Gulo_3000_Poster.pptx')

# Handout is created after the poster PDF has been exported and thumbnailed.
if '--poster-only' in sys.argv:
    print('Poster written.');sys.exit(0)

doc=Document();sec=doc.sections[0]
sec.page_width=DInches(8.5);sec.page_height=DInches(11)
sec.top_margin=DInches(.48);sec.bottom_margin=DInches(.43)
sec.left_margin=sec.right_margin=DInches(.52)
sec.header_distance=DInches(.16);sec.footer_distance=DInches(.17)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=DPt(10)
normal.paragraph_format.space_after=DPt(4);normal.paragraph_format.line_spacing=1.02
for name,sz in [('Title',18),('Heading 1',12),('Heading 2',10)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=DPt(sz);st.font.bold=True;st.font.color.rgb=DColor.from_string(NAVY)
    st.paragraph_format.space_before=DPt(5);st.paragraph_format.space_after=DPt(3)
headerp=sec.header.paragraphs[0];headerp.text='POWERTRAIN DYNAMICS GROUP   |   ME 495   |   SEPTEMBER 2026'
headerp.runs[0].font.size=DPt(7);headerp.runs[0].font.color.rgb=DColor.from_string(GRAY)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Gulo 3000 • Technical handout   |   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
for r in fp.runs:r.font.size=DPt(7)

def para(s,style=None):
    return doc.add_paragraph(s,style)
def head(s):return doc.add_paragraph(s,'Heading 1')
def pic(name,width=7.40):
    p=doc.add_paragraph();p.paragraph_format.space_after=DPt(1)
    p.add_run().add_picture(str(FIG/(name+'.png')),width=DInches(width));return p
def caption(s):
    p=para(s)
    for r in p.runs:r.font.size=DPt(8);r.font.color.rgb=DColor.from_string(GRAY)
    return p
def dtable(rows,widths=None,fontsize=8):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    if widths:
        for col,wd in zip(t.columns,widths):col.width=DInches(wd)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        if widths:
            for cell,wd in zip(cells,widths):cell.width=DInches(wd)
        for c,s in zip(cells,row):
            c.text=str(s)
            for p in c.paragraphs:
                p.paragraph_format.space_after=DPt(2);p.paragraph_format.space_before=DPt(1)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs:
                    r.font.name='Arial';r.font.size=DPt(fontsize);r.font.bold=i==0
                    r.font.color.rgb=DColor.from_string(WHITE if i==0 else NAVY)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),NAVY if i==0 else (PALE if i%2 else WHITE));c._tc.get_or_add_tcPr().append(sh)
        trpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    return t

para('Reducing Gulo 3000 shaft vibration','Title')
para('1 / Experimental basis and mathematical model','Heading 1')
thumb=OUT/'Presentation_Source/poster_thumbnail.png'
t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=DInches(3.55);t.columns[1].width=DInches(3.85)
if thumb.exists():t.cell(0,0).paragraphs[0].add_run().add_picture(str(thumb),width=DInches(3.5))
p=t.cell(0,1).paragraphs[0]
p.add_run('Recommendation. ').bold=True
p.add_run(f'Investigate a shaft-path damper with Bs = {DES["Bs"]["Proposed"]:.3f} N·m·s/rad. It reduces modeled switching ripple by {DES["Bs"]["TransientReduction_pct"]:.1f}% and preserves steady-speed gain. Four alternatives meet the nominal 30% target under the stated fixed-voltage test. These are model predictions; no modified hardware was tested.')
for r in p.runs:r.font.size=DPt(9)

head('Data and experimental procedures')
para(f'The supplied September 10 records contain {V["sweepCount"]} distinct frequency sweeps ({V["rawSweepRows"]} rows, 4–10 Hz) and {V["timeRecordCount"]} time histories. The copied/renamed exports are byte-identical to the originals and are not counted twice. The LabVIEW headers already express both tachometers in rad/s. Gain is distal-speed amplitude divided by motor-voltage amplitude; repeated complex responses are averaged in 0.05-Hz bins ({V["frequencyBins"]} bins). Phase is treated circularly before unwrapping. [2,3]')
para('DC resistance uses the two stationary tests. Km is the slope of V − RmI versus steady motor speed across 14 records, with a free intercept. The 20-Hz recordings are retained as diagnostics rather than pooled as DC resistance. Locked-side free decays give Btr = Bs + Bb and Btl = Bs + Bm. For the two coastdowns, the motor was physically uncoupled while both flywheels remained connected by the shaft (confirmed by the team).')

head('Equations and parameter provenance')
para('States: shaft twist δ = θ1 − θ2, motor speed ω1, distal speed ω2 and armature current i. The model uses Jc = Jm + Jf and the same Jf for both flywheels. [2]')
eq=para('Lm di/dt = Vin − Rm i − Km ω1\nJc dω1/dt = Km i − Bm ω1 − Bs(ω1 − ω2) − Kδ\nJf dω2/dt = Bs(ω1 − ω2) + Kδ − Bb ω2;     dδ/dt = ω1 − ω2')
for r in eq.runs:r.font.name='Cambria';r.font.size=DPt(9)
para('The frequency model H(s) = Ω2(s)/Vin(s) neglects Lm; switching simulations retain it. Jf = ρπDf⁴Lf/32 and Kgeometry = πGsDs⁴/(32Ls). The complete third-order polynomial and an independent state-equation check are in the MATLAB functions.')
rows=[['Parameter','Value','Basis'],
 ['Rm; Km',f'{P["Rm"]:.4f} Ω; {P["Km"]:.5f} V·s/rad','Stationary tests; back-EMF regression'],
 ['Jf; Jm',f'{V["Jf"]:.6f}; {P["Jm"]:.3g} kg·m²','Given geometry/material; given rotor inertia'],
 ['Kgeometry → Keffective',f'{V["K_geometry"]:.4f} → {V["K_refined"]:.4f} N·m/rad','Geometry; only parameter fitted to sweeps'],
 ['Btr; Btl',f'{V["Btr"]:.6f}; {V["Btl"]:.6f} N·m·s/rad','Mean of two free decays per locked side'],
 ['Bb; Bs; Bm',f'{P["Bb"]:.7f}; {P["Bs"]:.6f}; {P["Bm"]:.6f} N·m·s/rad','Coastdown and subtraction of damping sums'],
 ['Lm; shaft geometry',f'{P["Lm"]:.3f} H; Ds = 3.18 mm; Ls = 305 mm','Given course parameters [2]']]
dtable(rows,[1.65,2.80,2.95],8.1)
head('Damping extraction and its limits')
para('Free-decay extrema are fitted as ln|ωpeak| = c − αt; B = 2Jα, ωn = √[(2πfd)² + α²] and ζ = α/ωn. Extrema above 15% of the initial amplitude, within 8 s, are retained. The coast fit uses 4.0–19.5 s after disengagement and J = 2Jf, giving Bb ≈ 2Jfα. Then Bs = Btr − Bb and Bm = Btl − Bs. All three are effective coefficients of this assembly, not independently isolated material properties.')
para(f'The motor regression has R² = {V["motorR2"]:.5f} and intercept {V["motorOffset"]:.3f} V. The intercept and effective resistance absorb offsets/brush effects. Coastdown is slightly better fit by a linear Coulomb model than an exponential model; an equivalent viscous model is retained for frequency-domain design. This approximation limits extrapolation.')

doc.add_page_break()
para('2 / Validation and design screening','Heading 1')
pic('model_validation',5.35)
caption('Figure 1. Measured response, flexible model and rigid-shaft limit. A rigid shaft enforces ω1 = ω2 and removes the torsional resonance; setting Bs = 0 does not make a shaft rigid.')
dtable([['Check','Magnitude / frequency result','Phase / interpretation'],
 ['Measured vs modeled resonance',f'{V["measuredResonance"]:.3f} vs {V["modelResonance"]:.3f} Hz',f'Peak gain {V["measuredPeak"]:.3f} vs {V["modelPeak"]:.3f} rad/(s·V)'],
 ['Before → after K refinement',f'{V["initialMagnitudeRMSE"]:.2f} → {V["refinedMagnitudeRMSE"]:.2f} dB RMS',f'{V["initialPhaseRMSE"]:.2f}° → {V["refinedPhaseRMSE"]:.2f}° RMS'],
 ['Leave one sweep out','0.89–1.83 dB RMS','5.03–17.47° RMS; shared apparatus'],
 ['Broadband last-test diagnostic',f'{V["noiseValidation"]["magnitudeRMSE_dB"]:.2f} dB RMS',f'{V["noiseValidation"]["phaseRMSE_deg"]:.2f}° RMS; coherence ≥0.8']],
 [2.0,2.65,2.75],8.0)
para(f'K is refined from geometry by fitting mean squared log-magnitude and wrapped-phase residuals. It is an effective stiffness correction, not a new measured shear modulus. Near resonance (5.2–6.2 Hz), magnitude RMS error is {V["resonanceRegionMagnitudeRMSE"]:.2f} dB. Including the given inductance changes predictions by at most {V["maxInductanceDifference_dB"]:.3f} dB and {V["maxInductanceDifference_deg"]:.2f}° in the measured band.')
para('The broadband record is withheld from fitting: an H1 spectral estimate uses nine 2-s Hann windows with 50% overlap and 0.5-Hz resolution. Nine 4–10 Hz bins meet coherence ≥0.8. Its coarse resolution cannot resolve the narrow peak accurately. The three sweeps share setup and nearly constant excitation (~1.72–1.76 V amplitude); no descending sweep or systematic amplitude sweep establishes broad linearity. The model supports comparative design near the tested range, with residual discrepancies retained.')
head('Sensitivity identifies four attenuation directions')
para('Course cost: J = ∫|H(jω)|dω over ωr ± π (a 1-Hz-wide band); integrate linear magnitude, not dB. Normalized sensitivity is Sp = (ΔJ/J)/(Δp/p), estimated with symmetric ±1% perturbations. Record both the original-band cost and the cost following the new resonance, plus peak-height sensitivity. Recompute linked parameters after each perturbation. [2]')
sens={s['Parameter']:s for s in V['sensitivity']}
rows=[['Parameter','Original-band cost S','Tracked-peak S','Interpretation']]
for n in ['Rm','Km','Bm','Bs','Df','Ds']:
    rows.append([n,f'{sens[n]["FixedBandCostSensitivity"]:+.3f}',f'{sens[n]["PeakSensitivity"]:+.3f}',
                 'Attenuation candidate' if n in DES else 'Primarily moves resonance'])
dtable(rows,[1.05,1.85,1.65,2.85],8)
para('The four largest absolute peak-height sensitivities are Rm, Km, Bm and Bs. All 12 raw parameters were screened; Jm, Lf, ρf, Ls, Gs and Bb are also included in sensitivity.csv. Changing Df changes both flywheel inertias as Df⁴; Ds changes stiffness as Ds⁴. Geometry can move the resonance out of a fixed cost band while leaving its height nearly unchanged, so it is not selected on fixed-band improvement alone.')

doc.add_page_break()
para('3 / Four design alternatives and recommendation','Heading 1')
para('Targets below include a modest parameter margin beyond the minimum joint solution and are rounded outward to two significant digits. Each nominal design reduces both the frequency-response criterion and the defined switching ripple by at least 30%. These are component-level targets for prototype evaluation, not validated vehicle designs.')
rows=[['Parameter change','Gain ↓','1-Hz cost ↓','Pulse ↓','Speed gain retained']]
for n in ['Rm','Km','Bs','Bm']:
    d=DES[n]
    targ={'Rm':f'Rm: {P["Rm"]:.3f} → {d["Proposed"]:.2f} Ω',
          'Km':f'Km: {P["Km"]:.4f} → {d["Proposed"]:.3f} V·s/rad',
          'Bs':f'Bs: {P["Bs"]:.5f} → {d["Proposed"]:.3f} N·m·s/rad',
          'Bm':f'Bm: {P["Bm"]:.5f} → {d["Proposed"]:.3f} N·m·s/rad'}[n]
    rows.append([targ,percent(d['PeakReduction_pct']),percent(d['MovingCostReduction_pct']),percent(d['TransientReduction_pct']),percent(d['DCGainRatio']*100)])
dtable(rows,[3.0,1.0,1.1,1.0,1.30],8.1)
caption('For Bs the local peak is eliminated: “gain” is the maximum across 4–10 Hz, and the 1-Hz cost stays centered on the baseline resonance. For other designs, both metrics track their shifted local peak. Broadband low-frequency commanded rotation is not relabeled as torsional resonance.')
pic('switching_transient',5.35)
caption('Figure 2. Baseline and shaft-damper switching responses. Exact matrix-exponential simulation retains motor inductance; +1 V for 4 s, then a driven 0-V source for 4 s. Initial state is rest. Plots show the first 0.8 s after each switch.')
para(f'Torsional ripple is r = ω2 − (Jcω1 + Jfω2)/(Jc + Jf). The criterion is max|r| over the full pulse: baseline {V["baselineRipple"]:.4f} rad/s. It measures relative torsional motion, not total wheel speed or passenger acceleration. An open-circuit turn-off variant is also saved. The design gain/cost thresholds alone would need much smaller damping changes (see resonance_only_designs.csv) but would not achieve the pulse target.')
head('Practical comparison and parameter linkage')
para(f'1. Shaft-path damper — preferred concept. Add relative-motion dissipation across the shaft/coupling path, targeting Bs = {DES["Bs"]["Proposed"]:.3f} N·m·s/rad. This is about {DES["Bs"]["Factor"]:.1f} times the baseline coefficient and retains 100% of modeled steady-speed gain. Specify the damper by its measured torque–relative-speed behavior. Added stiffness, inertia, heat and packaging must be included in a new prototype model.')
para(f'2. Series resistance. Target total Rm = {DES["Rm"]["Proposed"]:.2f} Ω (about 1.61 Ω added). It is a simple passive concept, but produces I²R heat and retains only {DES["Rm"]["DCGainRatio"]*100:.0f}% of steady-speed gain. The input is then defined upstream of the added resistor; do not use downstream motor-terminal voltage as the same transfer-function input.')
para(f'3. Lower-Km motor. Target Km = {DES["Km"]["Proposed"]:.3f} V·s/rad (also N·m/A in SI). The idealized study holds Rm, Jm and damping fixed; real motor replacement changes these together. Refit all coupled properties and verify torque/current requirements. Modeled steady-speed gain retained: {DES["Km"]["DCGainRatio"]*100:.0f}%.')
para(f'4. Motor-side damper. Target Bm = {DES["Bm"]["Proposed"]:.3f} N·m·s/rad. This resists absolute motor speed and causes continuous dissipation. Only {DES["Bm"]["DCGainRatio"]*100:.0f}% of steady-speed gain remains; this makes it the least attractive of the four despite meeting the fixed-voltage target.')
para('Fixed input is a material limitation. If voltage is rescaled to restore equal final speed, pulse reductions are approximately '+
     ', '.join(f'{n}: {DES[n]["EqualSpeedTransientReduction_pct"]:.1f}%' for n in ['Rm','Km','Bs','Bm'])+
     '. A negative reduction means worse ripple. These rescaled simulations are diagnostic and are not recommended drive voltages. Eight combinations of the repeated damping tests are checked in replicate_robustness.csv; they are not statistical confidence intervals.')
head('References and reproduction')
caption('[1] Meyhofer & Reddy, Task Letter Lab 1, Sept. 4, 2026. [2] ME 495 Lab 1_1, pp. 12–25; Lab 1_2, pp. 4, 19–36, 43. [3] Flexible Shaft — Student Guide, pp. 20–27. Poster format follows Dr. Royston, F26 ME495 Posters-1, pp. 20–28, 48, 54–55. All numerical results: supplied Data/ exports; run MATLAB/run_all.m. Tables, diagnostic plots and verification results are in Analysis/; publication figures are in Figures/.')
doc.save(OUT/'Gulo_3000_Handout.docx')
print('Poster and handout source documents written.')
